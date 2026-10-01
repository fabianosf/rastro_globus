"""Helpers para importacao da planilha BASE GERAL (HistoricoGarantiaMensal)."""

from __future__ import annotations

import csv
import hashlib
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .models import Fornecedor, HistoricoGarantiaMensal

EMPRESAS = {c.value for c in HistoricoGarantiaMensal.Empresa}
SHEET_NAME = "BASE GERAL"

FORNECEDOR_ALIASES = {
    "BB TECH": "BBTECH",
}

HEADER_MAP = {
    "FORNECEDOR": "fornecedor",
    "SOLICITADO": "solicitado",
    "CONCEDIDO": "concedido",
    "EM ANALISE": "em_analise",
    "NEGADO": "negado",
    "EMPRESA": "empresa",
    "DATA": "data",
    "ANO": "ano",
    "MES": "mes",
}


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.upper()
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def normalize_fornecedor_nome(value: Any) -> str:
    norm = normalize_text(value)
    return FORNECEDOR_ALIASES.get(norm, norm)


def normalize_empresa(value: Any) -> str:
    return normalize_text(value)


def to_decimal(value: Any) -> Decimal | None:
    if value is None or value == "":
        return Decimal("0")
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).strip()
    if not text:
        return Decimal("0")
    text = text.replace("R$", "").replace(" ", "")
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return None


def to_competencia(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return date(value.year, value.month, 1)
    if isinstance(value, date):
        return date(value.year, value.month, 1)
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            dt = datetime.strptime(text[:10], fmt)
            return date(dt.year, dt.month, 1)
        except ValueError:
            continue
    return None


def hash_linha(
    norm_forn: str,
    empresa: str,
    competencia: date,
    concedido: Decimal,
    em_analise: Decimal,
    negado: Decimal,
) -> str:
    payload = (
        f"{norm_forn}|{empresa}|{competencia.isoformat()}|"
        f"{concedido:.2f}|{em_analise:.2f}|{negado:.2f}"
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def synthetic_cnpj(norm_forn: str) -> str:
    digest = hashlib.sha1(norm_forn.encode("utf-8")).hexdigest()[:14]
    return f"PL-{digest}"


def find_similar_pairs(names: list[str], threshold: float = 0.85) -> list[tuple[str, str, float]]:
    unique = sorted(set(n for n in names if n))
    pairs: list[tuple[str, str, float]] = []
    for i, a in enumerate(unique):
        for b in unique[i + 1 :]:
            ratio = SequenceMatcher(None, a, b).ratio()
            if ratio > threshold:
                pairs.append((a, b, round(ratio, 3)))
    pairs.sort(key=lambda x: (-x[2], x[0], x[1]))
    return pairs


def _map_headers(row: tuple[Any, ...]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(row):
        key = HEADER_MAP.get(normalize_text(cell))
        if key:
            mapping[key] = idx
    return mapping


def _cell(row: tuple[Any, ...], mapping: dict[str, int], key: str) -> Any:
    idx = mapping.get(key)
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _build_fornecedor_index() -> dict[str, Fornecedor]:
    index: dict[str, Fornecedor] = {}
    for forn in Fornecedor.objects.all().only(
        "id", "razao_social", "nome_fantasia", "cnpj", "origem"
    ):
        n1 = normalize_fornecedor_nome(forn.razao_social)
        if n1 and n1 not in index:
            index[n1] = forn
        if forn.nome_fantasia:
            n2 = normalize_fornecedor_nome(forn.nome_fantasia)
            if n2 and n2 not in index:
                index[n2] = forn
    return index


@dataclass
class ImportResult:
    inseridas: int = 0
    ignoradas: int = 0
    rejeitadas: int = 0
    fornecedores_criados: int = 0
    similares: list[tuple[str, str, float]] = field(default_factory=list)
    rejeitados_path: str = ""
    similares_path: str = ""


def import_planilha(caminho: Path, dry_run: bool = False) -> ImportResult:
    result = ImportResult()
    caminho = Path(caminho)
    suffix = caminho.suffix.lower()
    if suffix == ".csv":
        raise ValueError("Arquivo CSV nao e aceito. Use .xlsx da aba BASE GERAL.")
    if suffix not in {".xlsx", ".xls"}:
        raise ValueError(f"Extensao nao suportada: {suffix or '(vazia)'}. Use .xlsx.")
    if not caminho.is_file():
        raise FileNotFoundError(f"Arquivo nao encontrado: {caminho}")

    wb = load_workbook(caminho, read_only=True, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        wb.close()
        raise ValueError(f'Aba "{SHEET_NAME}" nao encontrada. Abas: {", ".join(wb.sheetnames)}')

    ws = wb[SHEET_NAME]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        wb.close()
        raise ValueError("Planilha vazia.")

    mapping = _map_headers(tuple(header))
    required = {"fornecedor", "concedido", "em_analise", "negado", "empresa", "data"}
    missing = required - set(mapping)
    if missing:
        wb.close()
        raise ValueError(f"Colunas obrigatorias ausentes: {', '.join(sorted(missing))}")

    rejeitados: list[list[Any]] = []
    nomes_norm: list[str] = []
    existing_hashes = set(
        HistoricoGarantiaMensal.objects.values_list("hash_linha", flat=True)
    )
    forn_index = _build_fornecedor_index()

    for row_num, row in enumerate(rows_iter, start=2):
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue

        nome_raw = _cell(row, mapping, "fornecedor")
        empresa_raw = _cell(row, mapping, "empresa")
        data_raw = _cell(row, mapping, "data")
        concedido = to_decimal(_cell(row, mapping, "concedido"))
        em_analise = to_decimal(_cell(row, mapping, "em_analise"))
        negado = to_decimal(_cell(row, mapping, "negado"))

        reasons: list[str] = []
        norm_forn = normalize_fornecedor_nome(nome_raw)
        if not norm_forn:
            reasons.append("fornecedor vazio")
        empresa = normalize_empresa(empresa_raw)
        if empresa not in EMPRESAS:
            reasons.append(f"empresa invalida: {empresa_raw!r}")
        competencia = to_competencia(data_raw)
        if competencia is None:
            reasons.append(f"data invalida: {data_raw!r}")
        if concedido is None:
            reasons.append("concedido invalido")
        elif concedido < 0:
            reasons.append("concedido < 0")
        if em_analise is None:
            reasons.append("em_analise invalido")
        elif em_analise < 0:
            reasons.append("em_analise < 0")
        if negado is None:
            reasons.append("negado invalido")
        elif negado < 0:
            reasons.append("negado < 0")

        if reasons:
            result.rejeitadas += 1
            rejeitados.append(
                [
                    row_num,
                    nome_raw,
                    empresa_raw,
                    data_raw,
                    _cell(row, mapping, "concedido"),
                    _cell(row, mapping, "em_analise"),
                    _cell(row, mapping, "negado"),
                    "; ".join(reasons),
                ]
            )
            continue

        assert competencia is not None and concedido is not None
        assert em_analise is not None and negado is not None

        solicitado = concedido + em_analise + negado
        h = hash_linha(norm_forn, empresa, competencia, concedido, em_analise, negado)
        nomes_norm.append(norm_forn)

        if h in existing_hashes:
            result.ignoradas += 1
            continue

        if dry_run:
            if norm_forn not in forn_index:
                result.fornecedores_criados += 1
                forn_index[norm_forn] = None  # type: ignore[assignment]
            existing_hashes.add(h)
            result.inseridas += 1
            continue

        forn = forn_index.get(norm_forn)
        if forn is None:
            forn = Fornecedor.objects.create(
                razao_social=str(nome_raw).strip() or norm_forn,
                nome_fantasia=str(nome_raw).strip() or norm_forn,
                cnpj=synthetic_cnpj(norm_forn),
                origem="planilha",
            )
            forn_index[norm_forn] = forn
            result.fornecedores_criados += 1

        HistoricoGarantiaMensal.objects.create(
            fornecedor=forn,
            empresa=empresa,
            competencia=competencia,
            valor_concedido=concedido,
            valor_em_analise=em_analise,
            valor_negado=negado,
            valor_solicitado=solicitado,
            origem="planilha",
            hash_linha=h,
        )
        existing_hashes.add(h)
        result.inseridas += 1

    wb.close()

    result.similares = find_similar_pairs(nomes_norm)

    if dry_run:
        return result

    if rejeitados:
        rejeitados_path = caminho.with_name(f"{caminho.stem}_rejeitados.csv")
        with rejeitados_path.open("w", encoding="utf-8-sig", newline="") as fh:
            writer = csv.writer(fh, delimiter=";")
            writer.writerow(
                [
                    "linha",
                    "fornecedor",
                    "empresa",
                    "data",
                    "concedido",
                    "em_analise",
                    "negado",
                    "motivo",
                ]
            )
            writer.writerows(rejeitados)
        result.rejeitados_path = str(rejeitados_path)

    if result.similares:
        similares_path = caminho.with_name(f"{caminho.stem}_similares.txt")
        lines = [f"{a} ~ {b} ({ratio})" for a, b, ratio in result.similares]
        similares_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
        result.similares_path = str(similares_path)

    return result
